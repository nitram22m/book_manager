"""Logica de negocio de Book Manager."""
from __future__ import annotations

import csv
from datetime import datetime, timezone
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import unicodedata
from typing import Callable, TypedDict

from openpyxl import Workbook
from openpyxl.styles import Font
import requests
from dotenv import load_dotenv

from book_manager.entities.entities import (
    CotizacionDolar,
    Libro,
    TipoCotizacion,
)
from book_manager.repositories.repositories import (
    SQLRepositorioBase,
    SQLRepositorioAuditoria,
    SQLRepositorioCotizacion,
    SQLRepositorioStock,
    crear_repositorio_autor,
    crear_repositorio_editorial,
    crear_repositorio_genero,
    crear_repositorio_libro,
    crear_repositorio_moneda,
    crear_repositorio_precio,
    crear_repositorio_tipo_cotizacion,
)
from book_manager.repositories.scraped_repository import (
    SQLRepositorioLibroScrapeado,
)
from book_manager.reportes import REPORTS_DIR
from book_manager.services.auditoria import auditar_operacion

ROOT_DIR: Path = Path(__file__).resolve().parents[3]
load_dotenv(ROOT_DIR / ".env")


def _ejecutar_scraper(
    query: str,
    max_items: int,
    queries: list[str] | None = None,
) -> list[dict[str, object]]:
    """Ejecuta Scrapy en un proceso aislado y devuelve el feed JSON."""
    with tempfile.TemporaryDirectory(
        prefix="book-manager-scrapy-"
    ) as directorio_temporal:
        ruta_feed = Path(directorio_temporal) / "resultados.json"
        entorno = os.environ.copy()
        ruta_src = str(ROOT_DIR / "src")
        pythonpath = entorno.get("PYTHONPATH", "")
        entorno["PYTHONPATH"] = os.pathsep.join(
            valor for valor in (ruta_src, pythonpath) if valor
        )
        comando = [
            sys.executable,
            "-m",
            "scrapy",
            "crawl",
            "cuspide_spider",
        ]
        if queries is None:
            comando.extend(("-a", f"query={query}"))
        else:
            comando.extend(
                (
                    "-a",
                    f"queries={json.dumps(queries, ensure_ascii=False)}",
                )
            )
        comando.extend(
            ("-a", f"max_items={max_items}", "-O", str(ruta_feed))
        )
        try:
            resultado = subprocess.run(
                comando,
                cwd=ROOT_DIR,
                env=entorno,
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
        except subprocess.TimeoutExpired as error:
            raise RuntimeError(
                "La búsqueda en Cúspide superó el límite de 180 segundos."
            ) from error

        if resultado.returncode != 0:
            detalle = resultado.stderr.strip()[-2000:]
            raise RuntimeError(
                "Falló la ejecución de Scrapy."
                + (f" Detalle: {detalle}" if detalle else "")
            )
        if not ruta_feed.is_file():
            raise RuntimeError("Scrapy terminó sin generar el feed JSON.")
        try:
            datos = json.loads(ruta_feed.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            raise RuntimeError(
                "No se pudo leer el feed JSON generado por Scrapy."
            ) from error

    if not isinstance(datos, list) or any(
        not isinstance(item, dict) for item in datos
    ):
        raise RuntimeError("El feed JSON de Scrapy tiene un formato inválido.")
    return datos


def _normalizar_titulo(titulo: str) -> str:
    """Normaliza títulos para comparar acentos, signos y capitalización."""
    sin_acentos = unicodedata.normalize("NFKD", titulo)
    caracteres = sin_acentos.encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", caracteres.casefold()).strip()


class AlertaPrecio(TypedDict):
    """Campos exportables de una diferencia de precio relevante."""

    titulo: str
    precio_local: float
    precio_cuspide: float
    diferencia: float
    moneda: str
    stock_cuspide: int
    url: str
    fecha_extraccion: datetime


class ServicioLibreria:
    """Conecta la interfaz de consola con los repositorios SQL.

    Es el unico lugar donde se arma el grafo de repositorios, de modo que
    cada servicio use el repositorio de persistencia correspondiente.
    """

    def __init__(
        self,
        scraper: Callable[
            [str, int, list[str] | None],
            list[dict[str, object]],
        ] | None = None,
    ) -> None:
        """Crea y conecta los repositorios necesarios para el servicio."""
        self.repo_generos = crear_repositorio_genero()
        self.repo_editoriales = crear_repositorio_editorial()
        self.repo_autores = crear_repositorio_autor()
        self.repo_monedas = crear_repositorio_moneda()
        self.repo_tipo_cot = crear_repositorio_tipo_cotizacion()
        self.repo_precios = crear_repositorio_precio()
        self.repo_libros = crear_repositorio_libro()
        self.repo_stock = SQLRepositorioStock()
        self.repo_cotizaciones = SQLRepositorioCotizacion()
        self.repo_auditoria = SQLRepositorioAuditoria()
        self.repo_libros_scrapeados = SQLRepositorioLibroScrapeado()
        self.servicio_cotizaciones = ServicioCotizacionDolar(
            self.repo_tipo_cot,
            self.repo_cotizaciones,
        )
        self.__scraper = scraper or _ejecutar_scraper

    @auditar_operacion
    def buscar_libros_cuspide(
        self,
        titulo: str,
        umbral: float,
        ruta_salida: Path | None = None,
    ) -> tuple[
        list[tuple[dict[str, object], Libro | None, int | None]],
        list[AlertaPrecio],
        Path,
        Path,
    ]:
        """Busca un libro, compara precios y exporta un reporte por ejecución.

        El reporte se genera aun cuando no haya diferencias sobre el umbral.
        """
        self._validar_umbral(umbral)
        libro, resultados_vinculados = (
            self._obtener_resultados_libro_cuspide(titulo)
        )
        resultados_scrapeados = [
            item for item, _, _ in resultados_vinculados
        ]
        alertas = self._comparar_resultados(
            [libro],
            resultados_scrapeados,
            umbral,
        )
        rutas = self._exportar_alertas(alertas, ruta_salida)
        return resultados_vinculados, alertas, rutas[0], rutas[1]

    def _obtener_resultados_libro_cuspide(
        self,
        titulo: str,
    ) -> tuple[
        Libro,
        list[tuple[dict[str, object], Libro | None, int | None]],
    ]:
        """Valida un libro local, lo scrapea y vincula resultados exactos."""
        consulta = titulo.strip()
        if not consulta:
            raise ValueError("El título de búsqueda no puede estar vacío.")

        libros_locales = self.repo_libros.leer_todos()
        libro_solicitado = self._buscar_libro_local(
            consulta,
            libros_locales,
        )
        if libro_solicitado is None:
            raise ValueError(
                f"'{consulta}' no pertenece al catálogo local. "
                "Solo se pueden buscar libros registrados."
            )

        resultados = self.__scraper(libro_solicitado.titulo, 10, None)
        resultados_vinculados: list[
            tuple[dict[str, object], Libro | None, int | None]
        ] = []
        for resultado in resultados[:10]:
            titulo_resultado = resultado.get("titulo")
            if not isinstance(titulo_resultado, str):
                continue
            libro_local = self._buscar_libro_local(
                titulo_resultado,
                [libro_solicitado],
            )
            if libro_local is None:
                continue
            stock_local = self.consultar_stock(libro_local.id)
            resultados_vinculados.append(
                (resultado, libro_local, stock_local)
            )
        return libro_solicitado, resultados_vinculados

    @auditar_operacion
    def comparar_precios_cuspide(
        self,
        umbral: float,
        ruta_salida: Path | None = None,
    ) -> tuple[list[AlertaPrecio], Path, Path]:
        """Compara precios y exporta CSV y Excel con las alertas.

        El umbral es una diferencia absoluta en la moneda del precio local.
        Para no comparar importes incompatibles, se omiten libros cuya moneda
        local no coincide con la publicada por Cúspide.
        """
        self._validar_umbral(umbral)

        libros_locales = self.repo_libros.leer_todos()
        resultados_scrapeados = (
            self.__scraper(
                "",
                10,
                [libro.titulo for libro in libros_locales],
            )
            if libros_locales
            else []
        )
        alertas = self._comparar_resultados(
            libros_locales,
            resultados_scrapeados,
            umbral,
        )
        rutas = self._exportar_alertas(alertas, ruta_salida)
        return alertas, rutas[0], rutas[1]

    @auditar_operacion
    def ejecutar_scraping_catalogo(
        self,
        umbral: float,
        ruta_salida: Path | None = None,
    ) -> tuple[list[AlertaPrecio], Path, Path]:
        """Ejecuta el spider para el catálogo y genera sus reportes."""
        return self.comparar_precios_cuspide(umbral, ruta_salida)

    @auditar_operacion
    def generar_reporte_precios_cuspide(
        self,
        umbral: float,
        ruta_salida: Path | None = None,
    ) -> tuple[list[AlertaPrecio], Path, Path]:
        """Genera reportes desde las observaciones guardadas, sin scrapear."""
        self._validar_umbral(umbral)
        libros_locales = self.repo_libros.leer_todos()
        resultados_scrapeados = self.repo_libros_scrapeados.leer_todos()
        alertas = self._comparar_resultados(
            libros_locales,
            resultados_scrapeados,
            umbral,
        )
        rutas = self._exportar_alertas(alertas, ruta_salida)
        return alertas, rutas[0], rutas[1]

    @auditar_operacion
    def leer_historial_auditoria(
        self,
        limite: int = 20,
    ) -> list[dict[str, object]]:
        """Devuelve los eventos de auditoría más recientes."""
        return self.repo_auditoria.leer_recientes(limite)

    @staticmethod
    def _validar_umbral(umbral: float) -> None:
        """Rechaza umbrales negativos, no numéricos o no finitos."""
        if (
            isinstance(umbral, bool)
            or not isinstance(umbral, (int, float))
            or not math.isfinite(umbral)
            or umbral < 0
        ):
            raise ValueError(
                "El umbral debe ser un número finito no negativo."
            )

    @staticmethod
    def _comparar_resultados(
        libros_locales: list[Libro],
        resultados_scrapeados: list[dict[str, object]],
        umbral: float,
    ) -> list[AlertaPrecio]:
        """Construye alertas para coincidencias con moneda compatible."""
        alertas: list[AlertaPrecio] = []
        for libro in libros_locales:
            coincidencia = ServicioLibreria._resultado_para_libro(
                libro.titulo,
                resultados_scrapeados,
            )
            if coincidencia is None:
                continue

            precio_externo = coincidencia.get("precio")
            moneda_externa = coincidencia.get("moneda")
            if (
                isinstance(precio_externo, bool)
                or not isinstance(precio_externo, (int, float))
                or not isinstance(moneda_externa, str)
            ):
                continue
            moneda_local = libro.precio.moneda.nombre
            if moneda_externa.upper() != moneda_local.upper():
                continue

            fecha_extraccion = ServicioLibreria._parsear_fecha_extraccion(
                coincidencia.get("fecha_actualizacion")
            )
            diferencia = abs(libro.precio.valor - float(precio_externo))
            if diferencia > umbral:
                stock_externo = coincidencia.get(
                    "disponibilidad_stock",
                    0,
                )
                url_externa = coincidencia.get("url", "")
                if (
                    isinstance(stock_externo, bool)
                    or not isinstance(stock_externo, int)
                    or not isinstance(url_externa, str)
                ):
                    continue
                alertas.append(
                    AlertaPrecio(
                        titulo=libro.titulo,
                        precio_local=libro.precio.valor,
                        precio_cuspide=float(precio_externo),
                        diferencia=diferencia,
                        moneda=moneda_local,
                        stock_cuspide=stock_externo,
                        url=url_externa,
                        fecha_extraccion=fecha_extraccion,
                    )
                )
        return alertas

    @staticmethod
    def _parsear_fecha_extraccion(valor: object) -> datetime:
        """Convierte la fecha del item a un datetime UTC con zona horaria."""
        if isinstance(valor, datetime):
            fecha = valor
        elif isinstance(valor, str):
            try:
                fecha = datetime.fromisoformat(
                    valor.replace("Z", "+00:00")
                )
            except ValueError as error:
                raise ValueError(
                    "La fecha de extracción del libro scrapeado no es válida."
                ) from error
        else:
            raise ValueError(
                "El resultado scrapeado no contiene fecha de extracción."
            )
        if fecha.tzinfo is None:
            fecha = fecha.replace(tzinfo=timezone.utc)
        return fecha.astimezone(timezone.utc)

    @staticmethod
    def _exportar_alertas(
        alertas: list[AlertaPrecio],
        ruta_salida: Path | None,
    ) -> tuple[Path, Path]:
        """Escribe CSV y XLSX con los datos requeridos para las alertas."""
        ruta_csv = ruta_salida or ServicioLibreria._ruta_reporte_alertas()
        if ruta_csv.suffix.lower() != ".csv":
            raise ValueError(
                "La ruta de salida del reporte debe terminar en .csv."
            )
        ruta_csv.parent.mkdir(parents=True, exist_ok=True)
        with ruta_csv.open(
            "w",
            newline="",
            encoding="utf-8-sig",
        ) as archivo:
            escritor = csv.writer(
                archivo,
            )
            escritor.writerow(
                (
                    "Libro",
                    "Precio interno",
                    "Precio web",
                    "Diferencia",
                    "Fecha de extracción",
                )
            )
            for alerta in alertas:
                escritor.writerow(
                    (
                        alerta["titulo"],
                        alerta["precio_local"],
                        alerta["precio_cuspide"],
                        alerta["diferencia"],
                        alerta["fecha_extraccion"].isoformat(),
                    )
                )
        ruta_excel = ruta_csv.with_suffix(".xlsx")
        ServicioLibreria._escribir_reporte_excel(alertas, ruta_excel)
        return ruta_csv, ruta_excel

    @staticmethod
    def _escribir_reporte_excel(
        alertas: list[AlertaPrecio],
        ruta_excel: Path,
    ) -> None:
        """Genera un libro Excel con las columnas pedidas en la consigna."""
        libro_excel = Workbook()
        hoja = libro_excel.active
        if hoja is None:
            raise RuntimeError("No se pudo crear la hoja del reporte Excel.")

        hoja.title = "Diferencias de precios"
        hoja.append(
            (
                "Libro",
                "Precio interno",
                "Precio web",
                "Diferencia",
                "Fecha de extracción",
            )
        )
        for celda in hoja[1]:
            celda.font = Font(bold=True)

        for alerta in alertas:
            hoja.append(
                (
                    alerta["titulo"],
                    alerta["precio_local"],
                    alerta["precio_cuspide"],
                    alerta["diferencia"],
                    alerta["fecha_extraccion"].replace(tzinfo=None),
                )
            )
            fila = hoja.max_row
            for columna in (2, 3, 4):
                hoja.cell(row=fila, column=columna).number_format = (
                    f'#,##0.00 "{alerta["moneda"]}"'
                )
            hoja.cell(row=fila, column=5).number_format = (
                "yyyy-mm-dd hh:mm:ss"
            )

        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = hoja.dimensions
        hoja.column_dimensions["A"].width = 38
        for columna in ("B", "C", "D"):
            hoja.column_dimensions[columna].width = 20
        hoja.column_dimensions["E"].width = 24
        libro_excel.save(ruta_excel)
        libro_excel.close()

    @staticmethod
    def _ruta_reporte_alertas() -> Path:
        """Genera la ruta CSV única de una ejecución."""
        marca_tiempo = datetime.now().astimezone().strftime(
            "%Y%m%d_%H%M%S_%f"
        )
        return REPORTS_DIR / (
            f"reporte_diferencias_precios_{marca_tiempo}.csv"
        )

    @staticmethod
    def _buscar_libro_local(
        titulo: str,
        libros: list[Libro],
    ) -> Libro | None:
        """Busca coincidencia exacta de título ignorando acentos."""
        titulo_normalizado = _normalizar_titulo(titulo)
        return next(
            (
                libro
                for libro in libros
                if _normalizar_titulo(libro.titulo) == titulo_normalizado
            ),
            None,
        )

    @staticmethod
    def _resultado_para_libro(
        titulo: str,
        resultados: list[dict[str, object]],
    ) -> dict[str, object] | None:
        """Selecciona el resultado scrapeado con el mismo título."""
        titulo_normalizado = _normalizar_titulo(titulo)
        return next(
            (
                resultado
                for resultado in resultados
                if isinstance(resultado.get("titulo"), str)
                and _normalizar_titulo(
                    str(resultado["titulo"])
                ) == titulo_normalizado
            ),
            None,
        )

    @auditar_operacion
    def calcular_precio_ars(
        self, libro_id: int, tipo_dolar_id: int = 1
    ) -> float:
        """Calcula el precio final del libro en ARS.

        Si el precio esta en USD se multiplica por la cotizacion mas reciente
        del tipo de dolar indicado.
        """
        libro = self.repo_libros.leer_por_id(libro_id)
        if libro is None:
            raise ValueError("El libro no existe.")

        precio = libro.precio
        moneda = precio.moneda.nombre
        if moneda == "ARS":
            return precio.valor
        if moneda != "USD":
            raise ValueError(f"No hay conversion configurada para {moneda}.")

        if self.repo_tipo_cot.leer_por_id(tipo_dolar_id) is None:
            raise ValueError("El tipo de cotizacion no existe.")
        historico = self.repo_cotizaciones.leer_historico_por_tipo(
            tipo_dolar_id
        )
        if not historico:
            raise ValueError(
                "No hay cotizaciones cargadas para hacer la conversión."
            )

        ultima = max(historico, key=lambda c: c.fecha)
        return precio.valor * ultima.valor

    @auditar_operacion
    def obtener_precios_bimonetarios(
        self,
        tipo_dolar_id: int,
        moneda_destino: str,
    ) -> tuple[
        list[tuple[Libro, float | None]],
        CotizacionDolar | None,
    ]:
        """Calcula precios en la moneda elegida y devuelve la cotización."""
        moneda_destino = moneda_destino.upper()
        if moneda_destino not in {"ARS", "USD"}:
            raise ValueError("La moneda destino debe ser ARS o USD.")
        if self.repo_tipo_cot.leer_por_id(tipo_dolar_id) is None:
            raise ValueError("El tipo de cotizacion no existe.")

        libros = self.repo_libros.leer_todos()
        necesita_cotizacion = any(
            libro.precio.moneda.nombre in {"ARS", "USD"}
            and libro.precio.moneda.nombre != moneda_destino
            for libro in libros
        )
        cotizacion_aplicada: CotizacionDolar | None = None
        if necesita_cotizacion:
            historico = self.repo_cotizaciones.leer_historico_por_tipo(
                tipo_dolar_id
            )
            if not historico:
                raise ValueError(
                    "No hay cotizaciones cargadas para hacer la conversión."
                )
            cotizacion_aplicada = max(
                historico,
                key=lambda registro: registro.fecha,
            )

        precios: list[tuple[Libro, float | None]] = []
        for libro in libros:
            precio = libro.precio
            moneda = precio.moneda.nombre
            if moneda == moneda_destino:
                precio_destino: float | None = precio.valor
            elif moneda in {"ARS", "USD"}:
                if cotizacion_aplicada is None:
                    raise ValueError(
                        "No hay cotización disponible para convertir precios."
                    )
                if moneda == "USD":
                    precio_destino = (
                        precio.valor * cotizacion_aplicada.valor
                    )
                else:
                    precio_destino = (
                        precio.valor / cotizacion_aplicada.valor
                    )
            else:
                precio_destino = None
            precios.append((libro, precio_destino))

        return precios, cotizacion_aplicada

    @auditar_operacion
    def obtener_precios_para_exportacion(
        self,
        tipo_dolar_id: int | None = None,
    ) -> list[tuple[Libro, float | None]]:
        """Devuelve el precio original y su equivalente ARS cuando aplica."""
        libros = self.repo_libros.leer_todos()
        necesita_cotizacion = any(
            libro.precio.moneda.nombre == "USD" for libro in libros
        )
        valor_cotizacion: float | None = None

        if necesita_cotizacion:
            if tipo_dolar_id is None:
                raise ValueError(
                    "Debe seleccionar un tipo de cotización "
                    "para convertir USD."
                )
            if self.repo_tipo_cot.leer_por_id(tipo_dolar_id) is None:
                raise ValueError("El tipo de cotizacion no existe.")
            historico = self.repo_cotizaciones.leer_historico_por_tipo(
                tipo_dolar_id
            )
            if not historico:
                raise ValueError(
                    "No hay cotizaciones cargadas para hacer la conversión."
                )
            valor_cotizacion = max(
                historico,
                key=lambda registro: registro.fecha,
            ).valor

        precios: list[tuple[Libro, float | None]] = []
        for libro in libros:
            precio = libro.precio
            moneda = precio.moneda.nombre
            if moneda == "ARS":
                equivalente_ars: float | None = precio.valor
            elif moneda == "USD":
                if valor_cotizacion is None:
                    raise ValueError(
                        "No hay cotización disponible para convertir USD."
                    )
                equivalente_ars = precio.valor * valor_cotizacion
            else:
                equivalente_ars = None
            precios.append((libro, equivalente_ars))

        return precios

    @auditar_operacion
    def consultar_stock(self, libro_id: int) -> int:
        """Devuelve las unidades disponibles, o cero si no hay registro."""
        stock = self.repo_stock.leer_por_libro(libro_id)
        return stock.cantidad if stock else 0

    @auditar_operacion
    def vender_libro(self, libro_id: int, cantidad_vendida: int) -> None:
        """Valida el libro y las existencias antes de descontar una venta."""
        if cantidad_vendida <= 0:
            raise ValueError("La cantidad a vender debe ser mayor que cero.")
        if self.repo_libros.leer_por_id(libro_id) is None:
            raise ValueError("El libro no existe.")
        stock = self.repo_stock.leer_por_libro(libro_id)
        if stock is None:
            raise ValueError("El libro no tiene stock registrado.")
        stock.descontar(cantidad_vendida)
        self.repo_stock.actualizar(stock)


class ServicioCotizacionDolar:
    """Obtiene y registra cotizaciones desde DolarApi."""

    _CASAS_A_TIPOS: dict[str, str] = {
        "oficial": "oficial",
        "blue": "blue",
        "bolsa": "mep",
        "contadoconliqui": "ccl",
        "mayorista": "mayorista",
        "cripto": "cripto",
        "tarjeta": "tarjeta",
    }
    _NOMBRES_TIPOS: dict[str, str] = {
        "oficial": "Oficial",
        "blue": "Blue",
        "mep": "MEP",
        "ccl": "CCL",
        "mayorista": "Mayorista",
        "cripto": "Cripto",
        "tarjeta": "Tarjeta",
    }

    def __init__(
        self,
        repositorio_tipos: SQLRepositorioBase[TipoCotizacion],
        repositorio_cotizaciones: SQLRepositorioCotizacion,
        api_url: str | None = None,
        solicitar: Callable[..., requests.Response] = requests.get,
    ) -> None:
        """Configura los repositorios y el acceso HTTP a la API."""
        self.__repositorio_tipos = repositorio_tipos
        self.__repositorio_cotizaciones = repositorio_cotizaciones
        self.__api_url = api_url or os.getenv("API_URL")
        if not self.__api_url:
            raise RuntimeError(
                "No se configuró API_URL. Defínala en el archivo .env."
            )
        self.__solicitar = solicitar

    @auditar_operacion
    def obtener_cotizaciones(self) -> list[CotizacionDolar]:
        """Descarga compra, venta y fecha de actualización desde DolarApi."""
        try:
            respuesta = self.__solicitar(self.__api_url, timeout=15)
            respuesta.raise_for_status()
        except requests.RequestException as error:
            raise RuntimeError(
                f"No se pudieron obtener cotizaciones desde {self.__api_url}."
            ) from error

        try:
            datos = respuesta.json()
        except ValueError as error:
            raise RuntimeError(
                "La API devolvió una respuesta JSON inválida."
            ) from error

        if not isinstance(datos, list):
            raise ValueError("La respuesta de la API debe ser una lista.")

        tipos = {
            tipo.nombre.casefold(): tipo
            for tipo in self.__repositorio_tipos.leer_todos()
        }
        siguiente_id = max(
            (tipo.id for tipo in tipos.values()),
            default=0,
        ) + 1
        cotizaciones: list[CotizacionDolar] = []
        tipos_procesados: set[str] = set()

        for indice, elemento in enumerate(datos, start=1):
            if not isinstance(elemento, dict):
                raise ValueError(
                    f"El registro {indice} de la API no es un objeto."
                )

            casa = elemento.get("casa")
            if not isinstance(casa, str):
                raise ValueError(
                    f"El registro {indice} de la API no contiene una casa."
                )

            nombre_tipo = self._CASAS_A_TIPOS.get(casa.casefold())
            if nombre_tipo is None:
                continue

            if nombre_tipo in tipos_procesados:
                raise ValueError(
                    f"La API devolvió más de una cotización para "
                    f"'{nombre_tipo}'."
                )

            fecha_actualizacion = elemento.get("fechaActualizacion")
            if not isinstance(fecha_actualizacion, str):
                raise ValueError(
                    f"El registro {indice} de la API no contiene "
                    "una fecha de actualización válida."
                )
            try:
                fecha_parseada = datetime.fromisoformat(
                    fecha_actualizacion.replace("Z", "+00:00")
                )
            except ValueError as error:
                raise ValueError(
                    f"La fecha de actualización del registro {indice} "
                    "no tiene un formato ISO válido."
                ) from error
            if fecha_parseada.tzinfo is None:
                fecha_parseada = fecha_parseada.replace(tzinfo=timezone.utc)
            fecha_cotizacion = fecha_parseada.astimezone(
                timezone.utc
            ).date()

            venta = elemento.get("venta")
            if (
                isinstance(venta, bool)
                or not isinstance(venta, (int, float))
                or not math.isfinite(venta)
                or venta <= 0
            ):
                raise ValueError(
                    f"La cotización de venta del registro {indice} "
                    "no es un número positivo válido."
                )

            compra = elemento.get("compra")
            if compra is not None and (
                isinstance(compra, bool)
                or not isinstance(compra, (int, float))
                or not math.isfinite(compra)
                or compra <= 0
            ):
                raise ValueError(
                    f"La cotización de compra del registro {indice} "
                    "no es un número positivo válido."
                )

            tipo = tipos.get(nombre_tipo)
            if tipo is None:
                tipo = TipoCotizacion(
                    siguiente_id,
                    self._NOMBRES_TIPOS[nombre_tipo],
                )
                self.__repositorio_tipos.crear(tipo)
                tipos[nombre_tipo] = tipo
                siguiente_id += 1

            cotizaciones.append(
                CotizacionDolar(
                    tipo,
                    fecha_cotizacion,
                    float(venta),
                    float(compra) if compra is not None else None,
                )
            )
            tipos_procesados.add(nombre_tipo)

        if not cotizaciones:
            raise ValueError(
                "La API no devolvió cotizaciones compatibles con el sistema."
            )

        return self.__repositorio_cotizaciones.guardar_o_actualizar_lote(
            cotizaciones
        )
